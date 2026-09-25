# -*- coding: utf-8 -*-
"""
EcoBin Connect presentation — Design System (DOC-DS) style
Includes: Architecture diagram, Wireframes, Atom → Molecule → Organism
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parent.parent
CAPTURES = REPO / "docs" / "design" / "captures"
ASSETS = ROOT / "assets"
OUT = ROOT / "EcoBin-Connect-Design-System.pptx"

# Palette matching Pattern library / UI docs
GREEN_DARK = RGBColor(0x16, 0x33, 0x2C)
GREEN = RGBColor(0x0D, 0x7A, 0x66)
GREEN_SOFT = RGBColor(0xED, 0xF5, 0xF1)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
PURPLE = RGBColor(0x5B, 0x4B, 0xB6)
INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x64, 0x74, 0x8B)
LINE = RGBColor(0xE2, 0xE8, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xFA, 0xFC, 0xFB)
NAVY = RGBColor(0x11, 0x18, 0x27)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
TOTAL = 12


def set_run(run, size=16, bold=False, color=INK, font="Segoe UI"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    node = f'<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{font}"/>'
    if ea is None:
        rPr.append(parse_xml(node))
    else:
        ea.set("typeface", font)


def write_block(shape, lines, default_size=16, default_color=INK):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.clear()
    first = True
    for item in lines:
        if isinstance(item, str):
            text, size, bold, color = item, default_size, False, default_color
        else:
            text = item[0]
            size = item[1] if len(item) > 1 else default_size
            bold = item[2] if len(item) > 2 else False
            color = item[3] if len(item) > 3 else default_color
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        if first:
            p.clear()
            first = False
        run = p.add_run()
        run.text = text
        set_run(run, size=size, bold=bold, color=color)
        p.space_after = Pt(6)
        p.space_before = Pt(0)


def rect(slide, l, t, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    return s


def round_rect(slide, l, t, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    return s


def textbox(slide, l, t, w, h):
    return slide.shapes.add_textbox(l, t, w, h)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def chrome(slide, active="wireframe"):
    """Sticky DOC-DS chrome: brand + 3 tabs."""
    rect(slide, 0, 0, SLIDE_W, Inches(0.72), WHITE, LINE)
    # brand
    tb = textbox(slide, Inches(0.4), Inches(0.18), Inches(3.2), Inches(0.4))
    write_block(tb, [("EcoBin PCRU", 16, True, GREEN_DARK)])
    # tabs
    tabs = [
        ("01 Wireframe", "wireframe"),
        ("02 Pattern library", "pattern"),
        ("03 UI", "ui"),
    ]
    x = Inches(4.2)
    for label, key in tabs:
        on = key == active
        fill = NAVY if on else WHITE
        fg = WHITE if on else MUTED
        border = NAVY if on else LINE
        box = round_rect(slide, x, Inches(0.16), Inches(2.35), Inches(0.42), fill, border)
        t = textbox(slide, x, Inches(0.22), Inches(2.35), Inches(0.35))
        write_block(t, [(label, 11, True, fg)])
        t.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        x += Inches(2.45)


def doc_head(slide, doc_id, title, lede, y=0.85):
    badge = round_rect(slide, Inches(0.4), Inches(y), Inches(1.55), Inches(0.32), GREEN_SOFT, GREEN)
    bt = textbox(slide, Inches(0.4), Inches(y + 0.02), Inches(1.55), Inches(0.28))
    write_block(bt, [(doc_id, 10, True, GREEN)])
    bt.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    tb = textbox(slide, Inches(2.1), Inches(y - 0.02), Inches(10.5), Inches(0.4))
    write_block(tb, [(title, 22, True, GREEN_DARK)])
    if lede:
        tb2 = textbox(slide, Inches(0.4), Inches(y + 0.4), Inches(12.4), Inches(0.35))
        write_block(tb2, [(lede, 12, False, MUTED)])


def footer(slide, page):
    rect(slide, 0, Inches(7.15), SLIDE_W, Inches(0.35), WHITE)
    rect(slide, 0, Inches(7.15), SLIDE_W, Pt(1), LINE)
    tb = textbox(slide, Inches(0.4), Inches(7.2), Inches(10), Inches(0.28))
    write_block(tb, [("docs/design  ·  EcoBin Connect  ·  PCRU", 10, False, MUTED)])
    tb2 = textbox(slide, Inches(11.6), Inches(7.2), Inches(1.4), Inches(0.28))
    write_block(tb2, [(f"{page} / {TOTAL}", 10, False, MUTED)])
    tb2.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT


def add_image(slide, path, left, top, width=None, height=None):
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)
    kwargs = {"left": left, "top": top}
    if width is not None:
        kwargs["width"] = width
    if height is not None:
        kwargs["height"] = height
    return slide.shapes.add_picture(str(path), **kwargs)


def picture_fit(slide, path, left, top, max_w, max_h):
    """Insert picture scaled to fit box, centered."""
    from PIL import Image

    path = Path(path)
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    w = iw * scale
    h = ih * scale
    # convert to EMU: PIL px assumed ~96dpi for pptx Inches math
    # pptx uses EMU; Inches() already; we pass float inches via Emu
    # easier: use Inches with computed inch sizes assuming 96 dpi
    w_in = (w / 96.0) if False else None
    # Use EMU directly from max box and aspect
    max_w_emu = int(max_w) if isinstance(max_w, int) else max_w
    # max_w/h are Inches objects (EmuLength)
    aspect = iw / ih
    box_aspect = max_w / max_h
    if aspect > box_aspect:
        pic_w = max_w
        pic_h = max_w / aspect
    else:
        pic_h = max_h
        pic_w = max_h * aspect
    cx = left + (max_w - pic_w) / 2
    cy = top + (max_h - pic_h) / 2
    return slide.shapes.add_picture(str(path), cx, cy, width=pic_w, height=pic_h)


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # ----- 1 Cover -----
    s = blank(prs)
    rect(s, 0, 0, SLIDE_W, SLIDE_H, GREEN_DARK)
    rect(s, 0, Inches(5.4), SLIDE_W, Inches(2.1), GREEN)
    tb = textbox(s, Inches(0.7), Inches(1.3), Inches(11.5), Inches(0.4))
    write_block(tb, [("DESIGN SYSTEM  ·  PRESENTATION", 13, True, GREEN_SOFT)])
    tb = textbox(s, Inches(0.7), Inches(1.9), Inches(11.5), Inches(2.2))
    write_block(
        s.shapes[-1] if False else tb,
        [
            ("EcoBin Connect", 44, True, WHITE),
            ("Architecture  ·  Wireframe  ·  Atomic Design", 20, False, GREEN_SOFT),
        ],
    )
    tb = textbox(s, Inches(0.7), Inches(4.3), Inches(11.5), Inches(0.8))
    write_block(
        tb,
        [
            ("DOC-DS-001 Wireframe  →  DOC-DS-002 Pattern  →  DOC-DS-003 UI", 14, False, GREEN_SOFT),
            ("มหาวิทยาลัยราชภัฏเพชรบูรณ์  ·  ปีการศึกษา 2568", 13, False, RGBColor(0xA8, 0xC4, 0xBA)),
        ],
    )
    tb = textbox(s, Inches(0.7), Inches(5.7), Inches(11.5), Inches(1.2))
    write_block(
        tb,
        [
            ("นายจิรกิตติ์ ตันตระกูล", 18, True, WHITE),
            ("อาจารย์ที่ปรึกษา  ผศ.ศรัญญา ตรีทศ", 13, False, WHITE),
        ],
    )

    # ----- 2 Agenda -----
    s = blank(prs)
    chrome(s, "pattern")
    doc_head(s, "AGENDA", "โครงเรื่องการนำเสนอ", "3 ชั้นของการออกแบบ ก่อนลงโค้ด")
    cards = [
        ("01", "Architecture", "สถาปัตยกรรมระบบจริง\nClient AI · Vercel · Cloud Run · SQL", GREEN_DARK),
        ("02", "Wireframe", "Low-fi B&W · โครงสร้างจอ\nMember / Guest / Admin", NAVY),
        ("03", "Atomic Design", "Atom → Molecule → Organism\nPattern library DOC-DS-002", PURPLE),
    ]
    for i, (num, title, desc, color) in enumerate(cards):
        x = Inches(0.45 + i * 4.2)
        round_rect(s, x, Inches(1.9), Inches(4.0), Inches(4.4), GREEN_SOFT, LINE)
        rect(s, x, Inches(1.9), Inches(4.0), Inches(0.9), color)
        nt = textbox(s, x, Inches(2.05), Inches(4.0), Inches(0.6))
        write_block(nt, [(f"{num}  {title}", 18, True, WHITE)])
        nt.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        dt = textbox(s, x + Inches(0.3), Inches(3.2), Inches(3.4), Inches(2.6))
        write_block(dt, [(desc, 15, False, INK)])
    footer(s, 2)

    # ----- 3 Architecture full diagram -----
    s = blank(prs)
    chrome(s, "ui")
    doc_head(s, "DOC-ARCH", "System Architecture", "ผู้ใช้ → Browser (UI + AI) → /api proxy → Go API → Cloud SQL")
    picture_fit(s, ASSETS / "architecture-diagram.png", Inches(0.35), Inches(1.55), Inches(12.6), Inches(5.4))
    footer(s, 3)

    # ----- 4 Architecture principles -----
    s = blank(prs)
    chrome(s, "ui")
    doc_head(s, "DOC-ARCH", "หลักการออกแบบสถาปัตยกรรม", "สรุปจาก docs/guides/02 + 06-uml-architecture")
    principles = [
        ("แยก Deploy", "Frontend / Backend / DB คนละที่ — ปิดโน้ตบุ๊กได้ ระบบยังรัน"),
        ("Client AI", "จำแนกขวดในเบราว์เซอร์ (MobileNet + EcoBin head) ไม่ส่งรูปไป classify ภายนอก"),
        ("Same-origin", "เบราว์เซอร์เรียก /api บนโดเมนเว็บ → Vercel proxy ไป Cloud Run"),
        ("แต้มทันที", "ผ่านเกณฑ์ ≥ ~80% → ให้แต้มทันที · Admin อ่านรายการสแกน"),
        ("Carbon EF", "คาร์บอน = มวล × Emission Factor (อ้าง TGO / CMH)"),
        ("3 บทบาท", "Member · Guest (ทดลอง ไม่ได้แต้มจริง) · Admin"),
    ]
    for i, (t, d) in enumerate(principles):
        x = Inches(0.4 + (i % 3) * 4.25)
        y = Inches(1.75 + (i // 3) * 2.45)
        round_rect(s, x, y, Inches(4.05), Inches(2.2), WHITE, LINE)
        rect(s, x, y, Inches(0.12), Inches(2.2), GREEN)
        tb = textbox(s, x + Inches(0.3), y + Inches(0.3), Inches(3.5), Inches(1.7))
        write_block(tb, [(t, 16, True, GREEN_DARK), (d, 13, False, INK)])
    footer(s, 4)

    # ----- 5 Wireframe overview -----
    s = blank(prs)
    chrome(s, "wireframe")
    doc_head(
        s,
        "DOC-DS-001",
        "Wireframe (Low-fidelity)",
        "ขาว–ดำ · โฟกัสโครงสร้างจอและ flow — ยังไม่ล็อกสี",
    )
    wf_set = [
        ("01 Login", CAPTURES / "wireframe" / "01-login.png"),
        ("02 Dashboard", CAPTURES / "wireframe" / "02-dashboard.png"),
        ("03 Scan", CAPTURES / "wireframe" / "03-scan.png"),
        ("05 Rewards", CAPTURES / "wireframe" / "05-rewards.png"),
        ("12 Admin", CAPTURES / "wireframe" / "12-admin-overview.png"),
        ("13 Admin QR", CAPTURES / "wireframe" / "13-admin-qr.png"),
    ]
    for i, (label, path) in enumerate(wf_set):
        x = Inches(0.35 + (i % 3) * 4.3)
        y = Inches(1.7 + (i // 3) * 2.6)
        round_rect(s, x, y, Inches(4.15), Inches(2.4), WHITE, LINE)
        caption = textbox(s, x + Inches(0.15), y + Inches(0.08), Inches(3.8), Inches(0.28))
        write_block(caption, [(label, 11, True, MUTED)])
        if path.exists():
            picture_fit(s, path, x + Inches(0.15), y + Inches(0.35), Inches(3.85), Inches(1.9))
    footer(s, 5)

    # ----- 6 Wireframe Member -----
    s = blank(prs)
    chrome(s, "wireframe")
    doc_head(s, "DOC-DS-001", "Wireframe — Member journey", "Login → Dashboard → Scan → Rewards / History")
    member_wf = [
        CAPTURES / "wireframe" / "01-login.png",
        CAPTURES / "wireframe" / "02-dashboard.png",
        CAPTURES / "wireframe" / "03-scan.png",
        CAPTURES / "wireframe" / "05-rewards.png",
    ]
    labels = ["1 เข้าสู่ระบบ", "2 แดชบอร์ด", "3 สแกนขยะ", "4 ของรางวัล"]
    for i, (path, label) in enumerate(zip(member_wf, labels)):
        x = Inches(0.3 + i * 3.25)
        round_rect(s, x, Inches(1.7), Inches(3.1), Inches(5.0), WHITE, LINE)
        cap = textbox(s, x + Inches(0.1), Inches(1.8), Inches(2.9), Inches(0.3))
        write_block(cap, [(label, 12, True, GREEN_DARK)])
        if path.exists():
            picture_fit(s, path, x + Inches(0.12), Inches(2.2), Inches(2.86), Inches(4.3))
    footer(s, 6)

    # ----- 7 Wireframe Admin -----
    s = blank(prs)
    chrome(s, "wireframe")
    doc_head(s, "DOC-DS-001", "Wireframe — Admin", "Overview · สแกน QR รับของ · รายการสแกน · จัดการรางวัล")
    admin_wf = [
        ("Overview", CAPTURES / "wireframe" / "12-admin-overview.png"),
        ("สแกน QR", CAPTURES / "wireframe" / "13-admin-qr.png"),
        ("รายการสแกน", CAPTURES / "wireframe" / "14-admin-scans.png"),
        ("ของรางวัล", CAPTURES / "wireframe" / "16-admin-rewards.png"),
    ]
    for i, (label, path) in enumerate(admin_wf):
        x = Inches(0.3 + i * 3.25)
        round_rect(s, x, Inches(1.7), Inches(3.1), Inches(5.0), WHITE, LINE)
        cap = textbox(s, x + Inches(0.1), Inches(1.8), Inches(2.9), Inches(0.3))
        write_block(cap, [(label, 12, True, PURPLE)])
        if path.exists():
            picture_fit(s, path, x + Inches(0.12), Inches(2.2), Inches(2.86), Inches(4.3))
    footer(s, 7)

    # ----- 8 Atomic hierarchy -----
    s = blank(prs)
    chrome(s, "pattern")
    doc_head(
        s,
        "DOC-DS-002",
        "Atomic Design — ลำดับชั้น",
        "Atom → Molecule → Organism  (Brad Frost) ใช้จัด Pattern library",
    )
    picture_fit(s, ASSETS / "atomic-hierarchy.png", Inches(0.5), Inches(1.55), Inches(12.3), Inches(5.4))
    footer(s, 8)

    # ----- 9 Atoms frame -----
    s = blank(prs)
    chrome(s, "pattern")
    doc_head(
        s,
        "DOC-DS-002",
        "ATOMS — พื้นฐาน UI",
        "สี · ฟอนต์ · ไอคอน · ปุ่ม · Input · Radio  + หลัก C.R.A.P.",
    )
    atoms = CAPTURES / "patterns" / "01-FRAME-1-ATOMS.png"
    if atoms.exists():
        picture_fit(s, atoms, Inches(0.4), Inches(1.55), Inches(12.5), Inches(5.4))
    footer(s, 9)

    # ----- 10 Molecules + Organisms frame -----
    s = blank(prs)
    chrome(s, "pattern")
    doc_head(
        s,
        "DOC-DS-002",
        "MOLECULES → ORGANISMS",
        "ชิ้นประกอบ → บล็อกหน้าจอที่ใช้จริงใน EcoBin",
    )
    mol = CAPTURES / "patterns" / "02-FRAME-2-MOLECULES-ORGANISMS-ECOBIN.png"
    if mol.exists():
        picture_fit(s, mol, Inches(0.35), Inches(1.5), Inches(12.6), Inches(5.45))
    footer(s, 10)

    # ----- 11 Selected components -----
    s = blank(prs)
    chrome(s, "pattern")
    doc_head(s, "DOC-DS-002", "ตัวอย่าง Organism สำคัญ", "Header · Scanner · Rewards · Redeem QR")
    picks = [
        ("App header", CAPTURES / "patterns" / "organism-16-Organism-App-header.png"),
        ("Scanner module", CAPTURES / "patterns" / "organism-19-Organism-Scanner-module.png"),
        ("Reward cards", CAPTURES / "patterns" / "organism-20-Organism-Reward-cards.png"),
        ("Redeem + QR", CAPTURES / "patterns" / "organism-22-Organism-Redeem-modal-QR.png"),
    ]
    for i, (label, path) in enumerate(picks):
        x = Inches(0.35 + (i % 2) * 6.45)
        y = Inches(1.65 + (i // 2) * 2.65)
        round_rect(s, x, y, Inches(6.25), Inches(2.45), WHITE, LINE)
        cap = textbox(s, x + Inches(0.2), y + Inches(0.1), Inches(5.8), Inches(0.28))
        write_block(cap, [(label, 12, True, GREEN_DARK)])
        if path.exists():
            picture_fit(s, path, x + Inches(0.2), y + Inches(0.4), Inches(5.85), Inches(1.9))
    footer(s, 11)

    # ----- 12 Fidelity ladder + close -----
    s = blank(prs)
    chrome(s, "ui")
    doc_head(s, "SUMMARY", "ลำดับ fidelity ที่ใช้ในโปรเจกต์", "ออกแบบทีละชั้น ก่อนลงโค้ด")
    steps = [
        ("01", "Wireframe", "Low-fi B&W\nโครงสร้าง + flow", NAVY, "DOC-DS-001"),
        ("02", "Pattern", "Atom → Organism\nคอมโพเนนต์กลาง", PURPLE, "DOC-DS-002"),
        ("03", "UI", "สี + โทเคน\nHigh-fi screens", GREEN, "DOC-DS-003"),
        ("04", "Prototype", "แอปบน Vercel\nใช้งานจริง", GREEN_DARK, "Production"),
    ]
    for i, (num, title, desc, color, doc) in enumerate(steps):
        x = Inches(0.4 + i * 3.2)
        round_rect(s, x, Inches(1.9), Inches(3.0), Inches(3.6), WHITE, LINE)
        rect(s, x, Inches(1.9), Inches(3.0), Inches(0.7), color)
        ht = textbox(s, x, Inches(2.0), Inches(3.0), Inches(0.5))
        write_block(ht, [(f"{num}  {title}", 14, True, WHITE)])
        ht.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        bt = textbox(s, x + Inches(0.2), Inches(2.9), Inches(2.6), Inches(1.8))
        write_block(bt, [(desc, 14, False, INK), ("", 8), (doc, 11, True, MUTED)])
        if i < 3:
            ar = textbox(s, x + Inches(2.85), Inches(3.4), Inches(0.4), Inches(0.4))
            write_block(ar, [("→", 20, True, AMBER)])
    note = textbox(s, Inches(0.5), Inches(5.8), Inches(12.2), Inches(1.0))
    write_block(
        note,
        [
            ("สรุป: Architecture บอกโครงสร้างระบบ · Wireframe บอกจอและ flow · Atomic Design บอกชิ้นประกอบที่ประกอบเป็น UI", 14, False, MUTED),
            ("ขอบคุณครับ — พร้อมรับคำถาม", 16, True, GREEN_DARK),
        ],
    )
    footer(s, 12)

    prs.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
